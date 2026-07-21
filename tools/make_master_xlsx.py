# -*- coding: utf-8 -*-
"""
特定技能 面談管理 マスタ.xlsx を生成する。
列名は Code.gs の参照名と完全一致させること（変更厳禁）。
デモデータ入り。Googleスプレッドシートに変換して使用する。
"""
import datetime
import secrets
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def tok():
    return secrets.token_hex(16)  # 32桁

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()

# ============================================================
# 1) 管理者マスタ
# ============================================================
ws = wb.active
ws.title = "管理者マスタ"
head = ["氏名", "メールアドレス（Googleアカウント）", "有効フラグ", "備考"]
ws.append(head)
rows = [
    ["髙橋 君成", "kiminari.takahashi@pluspivot.co.jp", "有効", "管理者（ホップ）"],
    ["（予備）ホップ管理者", "kimi1227nari@gmail.com", "有効", "予備アカウント"],
    ["退職者サンプル", "taishoku@example.com", "無効", "退職により無効化の例"],
]
for r in rows:
    ws.append(r)
set_widths(ws, [18, 40, 12, 26])
style_header(ws, len(head))

# ============================================================
# 2) 施設マスタ
#   面談周期(月) は半角カッコ！ PIN（任意・4桁） は全角カッコ！
# ============================================================
ws = wb.create_sheet("施設マスタ")
head = ["施設ID", "施設名", "所在地", "所属機関", "面談周期(月)",
        "アクセストークン", "PIN（任意・4桁）", "有効フラグ", "URLパラメータ（参考）"]
ws.append(head)
fac_rows = [
    ["F001", "さくら介護センター", "東京都新宿区西新宿1-2-3", "登録支援機関：株式会社ホップ", 3, tok(), "4821", "有効"],
    ["F002", "なごみ介護ホーム",   "神奈川県横浜市西区みなとみらい4-5", "登録支援機関：株式会社ホップ", 3, tok(), "",     "有効"],
    ["F003", "ひまわりケアタウン", "埼玉県さいたま市大宮区桜木町2-1", "登録支援機関：株式会社ホップ", 3, tok(), "1234", "有効"],
]
for i, r in enumerate(fac_rows, start=2):
    ws.append(r)
    # URLパラメータ（参考）：?fid=..&token=..(&pin=..) を自動生成する数式
    ws.cell(row=i, column=9).value = (
        '="?fid="&A{r}&"&token="&F{r}&IF(G{r}="","","&pin="&G{r})'.format(r=i)
    )
# PIN列を文字列書式に（先頭0対策）
for i in range(2, 2 + len(fac_rows)):
    ws.cell(row=i, column=7).number_format = "@"
set_widths(ws, [8, 20, 28, 26, 12, 34, 14, 10, 40])
style_header(ws, len(head))

# ============================================================
# 3) 人材マスタ  （列順は addWorker の appendRow と一致させること）
#   人材ID, 施設ID, 氏名（ローマ字）, 国籍, 在留資格, 配属日, 在籍状況, 備考
# ============================================================
ws = wb.create_sheet("人材マスタ")
head = ["人材ID", "施設ID", "氏名（ローマ字）", "国籍", "在留資格", "配属日", "在籍状況", "備考"]
ws.append(head)
d = datetime.date
worker_rows = [
    ["P0001", "F001", "TURNER ELIZABETH", "フィリピン", "特定技能1号", d(2024, 4, 1),  "在籍", "介護／居住費控除の是正事例あり"],
    ["P0002", "F001", "NGUYEN VAN A",     "ベトナム",   "特定技能1号", d(2024, 6, 15), "在籍", ""],
    ["P0003", "F002", "DELA CRUZ MARIA",  "フィリピン", "特定技能1号", d(2025, 1, 10), "在籍", ""],
    ["P0004", "F003", "LI WEI",           "中国",       "特定技能1号", d(2024, 10, 1), "在籍", "面談未実施（新規配属）"],
    ["P0005", "F002", "SHRESTHA RAM",     "ネパール",   "特定技能1号", d(2025, 3, 1),  "在籍", ""],
]
for r in worker_rows:
    ws.append(r)
set_widths(ws, [8, 8, 24, 12, 14, 14, 12, 30])
style_header(ws, len(head))

# ============================================================
# 4) 面談記録  （18項目コードは ITEMS と完全一致）
# ============================================================
ws = wb.create_sheet("面談記録")
ITEM_CODES = ["業務1","業務2","業務3","待遇1","待遇2","待遇3","待遇4_住居","待遇5_食費居住費","待遇6",
              "保護1","保護2","保護3","保護4","保護5","生活1","生活2","その他1","その他2"]
head = (["記録ID","人材ID","施設ID","面談日","対応者氏名","対応者役職"]
        + ITEM_CODES
        + ["問題内容","法令違反","その他特筆事項","次回面談予定日","作成年月日","入力者"])
ws.append(head)

def rec(recid, pid, fid, mdate, resp, role, problems, houki, note, other, nextdue, madeon, nyuryoku):
    """problems: {code: '有'} 指定分だけ有、他は無"""
    row = [recid, pid, fid, mdate, resp, role]
    for code in ITEM_CODES:
        row.append("有" if code in problems else "無")
    row += [note, houki, other, nextdue, madeon, nyuryoku]
    return row

records = [
    # TURNER：過去はOK → 直近で居住費控除の法令違反（期限切れ）
    rec("M00001","P0001","F001", d(2026,1,8), "支援 花子","支援担当者（主任）",
        {}, "無", "", "特になし。", d(2026,4,8), d(2026,1,8), "支援 花子"),
    rec("M00002","P0001","F001", d(2026,4,5), "支援 花子","支援担当者（主任）",
        {"待遇5_食費居住費":"有"}, "有",
        "②待遇 定期的に負担する食費，居住費等が合意したとおりの内容であること：居住費が合意額を超えて控除されていた（是正対応済み）",
        "特になし。", d(2026,7,5), d(2026,4,5), "支援 花子"),
    # NGUYEN：問題なし（次回はまだ先＝ok）
    rec("M00003","P0002","F001", d(2026,6,1), "支援 太郎","支援担当者",
        {}, "無", "", "特になし。", d(2026,9,1), d(2026,6,1), "支援 太郎"),
    # DELA CRUZ：健康面で問題あり（期限切れ）
    rec("M00004","P0003","F002", d(2026,4,15), "介護 次郎","施設長",
        {"生活2":"有"}, "無",
        "④生活 健康状態に異常がないこと：軽い体調不良の訴えあり。通院を案内し経過観察中。",
        "特になし。", d(2026,7,15), d(2026,4,15), "介護 次郎"),
    # SHRESTHA：問題なし（ok）
    rec("M00005","P0005","F002", d(2026,6,30), "介護 次郎","施設長",
        {}, "無", "", "特になし。", d(2026,9,30), d(2026,6,30), "介護 次郎"),
    # ※ P0004 LI WEI は面談記録なし＝「未実施」表示のデモ
]
for r in records:
    ws.append(r)
# 日付列の書式（面談日=4, 次回面談予定日=len-3? 明示指定）
date_cols_idx = [4, len(head)-2, len(head)-1]  # 面談日, 次回面談予定日, 作成年月日
# 正確に列位置を求める
col_map = {name: i+1 for i, name in enumerate(head)}
for i in range(2, 2 + len(records)):
    for name in ["面談日", "次回面談予定日", "作成年月日"]:
        ws.cell(row=i, column=col_map[name]).number_format = "yyyy/mm/dd"
set_widths(ws, [10,8,8,12,14,18] + [7]*len(ITEM_CODES) + [50,10,20,14,12,14])
style_header(ws, len(head))

# ============================================================
# 5) 法令違反記録 （列順は addRecord の appendRow と一致させること）
# ============================================================
ws = wb.create_sheet("法令違反記録")
head = ["違反ID","記録ID","人材ID","発生年月日","違反事実の内容",
        "ア_本人への対応","イ_責任者通知","イ_通知日","イ_通知先","イ_入管届出の案内",
        "ウ_関係機関通報","ウ_通報日","ウ_通報先"]
ws.append(head)
vio_rows = [
    ["V001","M00002","P0001", d(2026,3,25),
     "定期的に負担する居住費について、雇用契約・支援計画で合意した額を超える控除が行われていた（居住費の過大控除）。本人同意のない控除が確認されたため是正。",
     "労働基準監督署等の関係行政機関を案内", "通知済み", d(2026,4,6), "代表取締役・山田 太郎", "案内済み",
     "通報済み", d(2026,4,10), "東京出入国在留管理局"],
]
for r in vio_rows:
    ws.append(r)
for i in range(2, 2 + len(vio_rows)):
    for name in ["発生年月日","イ_通知日","ウ_通報日"]:
        ws.cell(row=i, column=head.index(name)+1).number_format = "yyyy/mm/dd"
set_widths(ws, [8,10,8,14,52,28,12,12,22,16,12,12,22])
style_header(ws, len(head))

# ============================================================
# 6) アクセスログ （ヘッダのみ。運用中は自動追記される）
# ============================================================
ws = wb.create_sheet("アクセスログ")
head = ["日時", "種別", "施設ID", "結果"]
ws.append(head)
set_widths(ws, [20, 12, 10, 24])
style_header(ws, len(head))

out = "/home/user/tokuteiginoreport/特定技能_面談管理_マスタ.xlsx"
wb.save(out)
print("saved:", out)
print("sheets:", wb.sheetnames)
